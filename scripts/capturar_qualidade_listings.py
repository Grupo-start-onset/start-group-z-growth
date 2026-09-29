# ==================================================================
# CAPTURA — Qualidade de Listings (Listings Items API) — EM LOTE
# Roda no Colab. Busca status e issues de cada ASIN via
# search_listings_items da Listings Items API, agrupando ate
# TAMANHO_LOTE ASINs por chamada (em vez de 1 por chamada), pra
# reduzir drasticamente o numero de chamadas e o tempo de execucao.
#
# Fonte de ASINs: ASINS_POR_CONTA_API.xlsx (uma aba por conta,
# lista real conferida manualmente por Lucas). Se a conta nao tem
# aba na planilha, cai no fallback (asins_catalogo_completo.json ->
# uniao de vendas_mes_*.json + estoque_mes_*.json).
#
# Balboa fica fora do mapeamento de aba por decisao do Lucas
# (25/09/2026): cache atual de 446 ASINs ja confirmado como correto,
# mesmo a planilha trazendo uma aba BALBOA com 350.
#
# Salva em: START_Vendor_Analytics/<conta>/raw/qualidade_listings.json
# Formato: { "<asin>": {"sku":..., "status":[...], "issues":[...]} }
#
# Tem cache por ASIN -- roda de novo sem problema, so busca o que
# falta. Faz limpeza: remove do cache ASINs que nao estao mais na
# lista real (planilha), pra nao arrastar ASIN morto/de outra conta.
#
# PADRAO daqui pra frente (decisao de 25/09/2026): capturas em lote
# na SP-API devem agrupar multiplos identificadores por chamada,
# nao uma chamada por item -- essa e a referencia a seguir pras
# proximas capturas (Catalog Items, etc.), nao so pra este script.
#
# CORRIGIDO em 25/09/2026: o parametro "identifiers" da API espera
# uma STRING separada por virgula (ex: "ASIN1,ASIN2,ASIN3"), nao uma
# lista Python. Passar uma lista nao dava erro -- so retornava 0
# resultados silenciosamente, o que fazia o script marcar todo mundo
# como "sem resultado no lote" (falso negativo). Corrigido com
# ",".join(lote) antes de montar a chamada.
# ==================================================================

import os, json, glob, time
from colab_shim import drive, userdata

try:
    from sp_api.api import ListingsItems
    from sp_api.base import Marketplaces
except ImportError:
    os.system("pip install python-amazon-sp-api -q")
    from sp_api.api import ListingsItems
    from sp_api.base import Marketplaces

try:
    import openpyxl
except ImportError:
    os.system("pip install openpyxl -q")
    import openpyxl

drive.mount('/content/drive', force_remount=True)

BASE_DRIVE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir, 'dados_raw'))
PLANILHA_ASINS = f'{BASE_DRIVE}/ASINS_POR_CONTA_API.xlsx'

CONTAS_CONFIG = {
    'alfa_jf':   {'pasta': 'alfa_jf/raw',    'secret': 'SP_API_REFRESH_TOKEN_ALFAJF'},
    'blidshop':  {'pasta': 'blidshop/raw',   'secret': 'SP_API_REFRESH_TOKEN_BLIDSHOP'},
    'conta3':    {'pasta': 'petclean/raw',   'secret': 'SP_API_REFRESH_TOKEN_PETCLEAN'},
    'ozitp':     {'pasta': 'ozitp/raw',      'secret': 'SP_API_REFRESH_TOKEN_OZITP'},
    'jolitex':   {'pasta': 'jolitex/raw',    'secret': 'SP_API_REFRESH_TOKEN_JOLITEX'},
    'balboa':    {'pasta': 'balboa/raw',     'secret': 'SP_API_REFRESH_TOKEN_BALBOA'},
    'riomaster': {'pasta': 'riomaster/raw',  'secret': 'SP_API_REFRESH_TOKEN_RIOMASTER'},
    'plastpet':  {'pasta': 'plastpet/raw',   'secret': 'SP_API_REFRESH_TOKEN_PLASTPET'},
}

SELLER_IDS = {
    'alfa_jf':   '76I78',
    'blidshop':  'UM8O1',
    'conta3':    'A470C',   # Petclean BR
    'ozitp':     'OZITV',
    'jolitex':   'R88OM',
    'balboa':    '6R8TT',
    'riomaster': 'RD8QP',
    'plastpet':  'SY933',   # Pet Factory Brazil Industria Ltda
}

# Nome da aba na planilha por conta (None = conta sem aba, usa fallback antigo)
ABA_PLANILHA = {
    'alfa_jf':   'ALFA',
    'blidshop':  'bLID',
    'conta3':    'PETCLEAN',
    'ozitp':     'ozitp',
    'jolitex':   'jOLITEX',
    'balboa':    None,      # mantido fora de proposito -- fica com cache atual (446)
    'riomaster': 'Rio Master',
    'plastpet':  None,      # sem aba na planilha -- usa fallback local
}

LWA_CLIENT_ID = userdata.get('SP_API_LWA_CLIENT_ID')
LWA_CLIENT_SECRET = userdata.get('SP_API_LWA_CLIENT_SECRET')
MARKETPLACE = Marketplaces.BR

TAMANHO_LOTE = 20       # ASINs por chamada
INTERVALO = 1.0         # pausa entre CHAMADAS (nao entre ASINs)
MAX_TENTATIVAS = 5

_planilha_cache = None

def carregar_planilha():
    global _planilha_cache
    if _planilha_cache is None:
        if os.path.exists(PLANILHA_ASINS):
            _planilha_cache = openpyxl.load_workbook(PLANILHA_ASINS, data_only=True)
            print(f'[planilha] carregada: {PLANILHA_ASINS} -- abas: {_planilha_cache.sheetnames}')
        else:
            _planilha_cache = False
            print(f'[aviso] planilha nao encontrada em {PLANILHA_ASINS}')
    return _planilha_cache


def asins_da_planilha(nome_aba):
    wb = carregar_planilha()
    if not wb or nome_aba not in wb.sheetnames:
        return None
    ws = wb[nome_aba]
    asins = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r and r[0]]
    return sorted(set(asins))


def asins_fallback(pasta_raw):
    caminho_completo = f'{pasta_raw}/asins_catalogo_completo.json'
    if os.path.exists(caminho_completo):
        try:
            lista = json.load(open(caminho_completo, encoding='utf-8'))
            if lista:
                return sorted(set(lista))
        except Exception:
            pass
    asins = set()
    for arq in glob.glob(f'{pasta_raw}/vendas_mes_*.json'):
        try:
            d = json.load(open(arq))
        except Exception:
            continue
        for r in d.get('salesByAsin', []):
            a = r.get('asin')
            if a:
                asins.add(a)
    for arq in glob.glob(f'{pasta_raw}/estoque_mes_*.json'):
        try:
            d = json.load(open(arq))
        except Exception:
            continue
        for r in d.get('inventoryByAsin', []):
            a = r.get('asin')
            if a:
                asins.add(a)
    return sorted(asins)


def asins_da_conta(chave, pasta_raw):
    nome_aba = ABA_PLANILHA.get(chave)
    if nome_aba:
        lista = asins_da_planilha(nome_aba)
        if lista:
            print(f'  [fonte: planilha aba "{nome_aba}"] {len(lista)} ASINs')
            return lista
        print(f'  [aviso] aba "{nome_aba}" nao encontrada/vazia -- fallback')
    lista = asins_fallback(pasta_raw)
    print(f'  [fonte: fallback local] {len(lista)} ASINs')
    return lista


def em_lotes(lista, tamanho):
    for i in range(0, len(lista), tamanho):
        yield lista[i:i + tamanho]


def buscar_lote_com_retry(api, seller_id, lote):
    identifiers_str = ",".join(lote)  # a API espera string separada por virgula, nao lista
    for tentativa in range(1, MAX_TENTATIVAS + 1):
        try:
            return api.search_listings_items(
                sellerId=seller_id,
                marketplaceIds=[MARKETPLACE.marketplace_id],
                identifiers=identifiers_str,
                identifiersType="ASIN",
                pageSize=len(lote),
                includedData=["summaries", "issues"],
            )
        except AttributeError:
            raise RuntimeError(
                "search_listings_items nao existe nesta versao da lib. "
                "Rode: !pip install --upgrade python-amazon-sp-api -q  e reinicie o runtime."
            )
        except Exception as e:
            texto_erro = str(e)
            if "QuotaExceeded" in texto_erro or "429" in texto_erro:
                espera = 2 ** tentativa
                print(f"    [cota excedida] aguardando {espera}s (tentativa {tentativa}/{MAX_TENTATIVAS})...")
                time.sleep(espera)
                continue
            raise
    raise RuntimeError("Numero maximo de tentativas excedido em search_listings_items")


def capturar_conta(chave, cfg):
    seller_id = SELLER_IDS.get(chave)
    if not seller_id:
        print(f'\n[pulando "{chave}"] SELLER_ID nao preenchido')
        return

    pasta_raw = os.path.join(BASE_DRIVE, cfg['pasta'])
    if not os.path.isdir(pasta_raw):
        print(f'\n[aviso] pasta nao encontrada para "{chave}": {pasta_raw} -- pulando')
        return

    print(f'\n=== {chave} ===')
    asins = asins_da_conta(chave, pasta_raw)
    if not asins:
        print('  nenhum ASIN encontrado -- pulando')
        return

    caminho_saida = f'{pasta_raw}/qualidade_listings.json'
    qualidade = {}
    if os.path.exists(caminho_saida):
        qualidade = json.load(open(caminho_saida, encoding='utf-8'))
        print(f'  {len(qualidade)} ja em cache de execucao anterior')

    removidos = [a for a in list(qualidade.keys()) if a not in asins]
    if removidos:
        print(f'  [limpeza] removendo {len(removidos)} ASINs que nao estao na lista real')
        for a in removidos:
            del qualidade[a]

    credentials = dict(
        refresh_token=userdata.get(cfg['secret']),
        lwa_app_id=LWA_CLIENT_ID,
        lwa_client_secret=LWA_CLIENT_SECRET,
    )
    api = ListingsItems(credentials=credentials, marketplace=MARKETPLACE)

    faltando = [a for a in asins if a not in qualidade or qualidade[a].get('erro')]
    print(f'  {len(faltando)} a buscar agora, em lotes de {TAMANHO_LOTE}')

    lotes = list(em_lotes(faltando, TAMANHO_LOTE))
    ok, falhas = 0, 0
    for i, lote in enumerate(lotes, 1):
        try:
            resp = buscar_lote_com_retry(api, seller_id, lote)
            items = (resp.payload or {}).get("items", [])
            encontrados = set()
            for item in items:
                asin_item = None
                if item.get("summaries"):
                    asin_item = item["summaries"][0].get("asin")
                if not asin_item:
                    continue
                encontrados.add(asin_item)
                qualidade[asin_item] = {
                    "sku": item.get("sku"),
                    "status": item.get("summaries", [{}])[0].get("status", []) if item.get("summaries") else [],
                    "issues": item.get("issues", []),
                }
                ok += 1
            # ASINs do lote que nao vieram na resposta
            for asin_faltante in set(lote) - encontrados:
                qualidade[asin_faltante] = {"sku": None, "status": [], "issues": [], "erro": "sem resultado no lote"}
                falhas += 1
        except RuntimeError as e:
            print(f'  [ERRO FATAL] {e}')
            break
        except Exception as e:
            for asin_erro in lote:
                qualidade[asin_erro] = {"sku": None, "status": [], "issues": [], "erro": str(e)}
            falhas += len(lote)
            print(f'  [erro no lote {i}] {str(e)[:80]}')

        if i % 5 == 0 or i == len(lotes):
            print(f'  [lote {i}/{len(lotes)}] ok={ok} falhas={falhas}')
            with open(caminho_saida, 'w', encoding='utf-8') as f:
                json.dump(qualidade, f, ensure_ascii=False, indent=2)
        time.sleep(INTERVALO)

    with open(caminho_saida, 'w', encoding='utf-8') as f:
        json.dump(qualidade, f, ensure_ascii=False, indent=2)
    print(f'  Salvo: {caminho_saida} ({len(qualidade)} ASINs, {ok} ok, {falhas} falhas)')


def main():
    for chave, cfg in CONTAS_CONFIG.items():
        capturar_conta(chave, cfg)
    print('\n\nConcluido.')


if __name__ == '__main__':
    main()
