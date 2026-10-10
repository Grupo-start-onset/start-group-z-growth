"""Gera as planilhas de controle de imagens (anúncio e A+) a partir dos listings ao vivo.

Uso (da raiz do repo):
  python3 scripts/imagens/gerar_planilha_controle.py <produtos.json> <saida_anuncio.xlsx> <saida_aplus.xlsx>

<produtos.json>: lista de produtos à venda (sku, asin, ean, nome, size, parent), extraída de
search_listings_items da conta, sem os SKUs "pai" de variação.
Depois de gerar, recalcular com o recalc.py do skill xlsx (as fórmulas saem sem valor em cache).
"""
import json, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.comments import Comment

F = 'Arial'
H_FILL = PatternFill('solid', fgColor='1D1D1F')
LOCK_FILL = PatternFill('solid', fgColor='EFEFF2')
IN_FILL = PatternFill('solid', fgColor='FFF2CC')
GREEN = PatternFill('solid', fgColor='D9F0E1')
RED = PatternFill('solid', fgColor='F8DADA')
ACC = 'E8711C'
thin = Side(style='thin', color='D0D0D6')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HR = 3  # linha do cabeçalho nas abas PRODUTOS e ARQUIVOS
DATA_CADASTRO = '10/10/2026'

# ------------------------------------------------------------------ conteúdo
GALERIA = [  # código, onde entra, o que mostrar (ração), (petiscos e higiene), obrigatória
    ('MAIN', 'Capa (posição 1)', 'Embalagem de frente, fundo branco puro, 85% da área, nada adicionado',
     'Embalagem de frente, fundo branco puro, 85% da área, nada adicionado', 'SIM'),
    ('PT01', 'Posição 2', 'Grande diferencial da linha (número ou promessa forte + grão e ingrediente)',
     'Grande diferencial (ex.: natural, biodegradável, perfumado)', 'SIM'),
    ('PT02', 'Posição 3', 'Benefícios com ícones (4 a 6, impressos na embalagem) + ingrediente',
     'Benefícios com ícones (4 a 6, impressos na embalagem)', 'SIM'),
    ('PT03', 'Posição 4', 'Formato e tamanho do grão, com medida em mm e referência de tamanho',
     'Detalhe e tamanho real (petisco, grânulo ou tapete aberto), com medidas', 'SIM'),
    ('PT04', 'Posição 5', 'Quantidade diária por peso do pet (números idênticos à embalagem)',
     'Modo de uso (como oferecer o petisco / como usar o produto)', 'SIM'),
    ('PT05', 'Posição 6', 'Transição alimentar dia a dia (igual ao verso da embalagem)',
     'Quantidade por dia ou rendimento do pacote', 'Recomendada'),
    ('PT06', 'Posição 7', 'Composição e níveis de garantia em letra grande',
     'Composição/níveis de garantia (petiscos) ou especificações (higiene)', 'SIM'),
    ('PT07', 'Posição 8', 'Emocional / momento de uso (pet do porte certo, pote cheio)',
     'Emocional / momento de uso', 'Recomendada'),
    ('PT08', 'Posição 9', 'Linha da marca ou embalagem anterior × nova',
     'Linha da marca ("conheça também")', 'Recomendada'),
]
APLUS = [  # código, (A: onde, ração, outros, tamanho), (B: onde, ração, outros, tamanho)
    ('AP01', ('Módulo 1 · Banner', 'Embalagem + pet + ingredientes, com o nome da linha', 'Produto + pet, com o nome da linha', '970 × 600'),
             ('Módulo 1 · Banner', 'Cena principal da linha', 'Cena principal da linha', '970 × 600')),
    ('AP02', ('Módulo 2 · Texto sobreposto', 'Principal diferencial; deixar o lado direito limpo', 'Principal diferencial; deixar o lado direito limpo', '970 × 300'),
             ('Módulo 2 · Imagem com destaques', 'Pacote ou grão em destaque', 'Produto em destaque', '300 × 300')),
    ('AP03', ('Módulo 3 · Benefício 1 de 4', 'Foto/ilustração do benefício 1', 'Foto/ilustração do benefício 1', '220 × 200'),
             ('Módulo 3 · Imagem 1 de 3', 'Ingrediente 1 (ex.: frango)', 'Benefício 1', '300 × 300')),
    ('AP04', ('Módulo 3 · Benefício 2 de 4', 'Foto/ilustração do benefício 2', 'Foto/ilustração do benefício 2', '220 × 200'),
             ('Módulo 3 · Imagem 2 de 3', 'Ingrediente 2 (ex.: arroz)', 'Benefício 2', '300 × 300')),
    ('AP05', ('Módulo 3 · Benefício 3 de 4', 'Foto/ilustração do benefício 3', 'Foto/ilustração do benefício 3', '220 × 200'),
             ('Módulo 3 · Imagem 3 de 3', 'Ingrediente 3 (ex.: linhaça)', 'Benefício 3', '300 × 300')),
    ('AP06', ('Módulo 3 · Benefício 4 de 4', 'Foto/ilustração do benefício 4', 'Foto/ilustração do benefício 4', '220 × 200'),
             ('Módulo 4 · Texto sobreposto', 'Momento de uso; deixar o lado direito limpo', 'Momento de uso; deixar o lado direito limpo', '970 × 300')),
    ('AP07', ('Módulo 4 · Especificações', 'Grão ou pacote em destaque', 'Produto em destaque', '300 × 300'),
             ('Não usado no modelo B', '-', '-', '-')),
    ('COMP', ('Módulo 5 · Tabela comparativa', 'Pacote na vertical, sem fundo colorido', 'Produto na vertical, sem fundo colorido', '150 × 300'),
             ('Módulo 5 · Tabela comparativa', 'Pacote na vertical, sem fundo colorido', 'Produto na vertical, sem fundo colorido', '150 × 300')),
]


# ------------------------------------------------------------------ dados
def tipo(nome):
    n = nome.lower()
    if 'ração' in n or 'sachê' in n:
        return 'Ração'
    if any(k in n for k in ('granulado', 'tapete', 'cata caca')):
        return 'Higiene'
    return 'Petisco'


def peso(nome, size):
    m = re.search(r'(\d+(?:[.,]\d+)?)\s*(kg|g)\b', nome, re.I)
    if m:
        return f"{m.group(1).replace('.', ',')} {m.group(2).lower()}"
    m = re.search(r'(\d+)\s*(un|unidades|refis)', nome, re.I)
    if m:
        return f'{m.group(1)} un.'
    return size or ''


def peso_kg(p):
    m = re.match(r'([\d,]+) (kg|g)', p)
    if not m:
        return 0
    v = float(m.group(1).replace(',', '.'))
    return v if m.group(2) == 'kg' else v / 1000


def preparar(rows):
    for r in rows:
        r['tipo'] = tipo(r['nome'])
        r['peso'] = peso(r['nome'], r['size'])
        r['familia'] = ''
        r['ref'] = ''
    fams = {}
    for r in rows:
        if r['parent']:
            fams.setdefault(r['parent'], []).append(r)
    for membros in fams.values():
        ref = max(membros, key=lambda r: peso_kg(r['peso']))  # A+ sugerido = maior embalagem
        base = re.sub(r'\s*-\s*Pacote.*$|\s*-\s*[\d,]+\s*kg$', '', ref['nome']).strip()
        for r in membros:
            r['familia'] = base
            r['ref'] = '' if r is ref else ref['ean']
    ordem = {'Ração': 0, 'Petisco': 1, 'Higiene': 2}
    rows.sort(key=lambda r: (ordem[r['tipo']], r['familia'] or 'zz' + r['nome'], -peso_kg(r['peso'])))
    return rows


# ------------------------------------------------------------------ helpers de estilo
def cabecalho(ws, row, headers, widths):
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row, c, h)
        cell.font = Font(name=F, size=10, bold=True, color='FFFFFF')
        cell.fill = H_FILL
        cell.alignment = Alignment(wrap_text=True, vertical='center')
        cell.border = BORDER
        ws.column_dimensions[cell.column_letter].width = widths[c - 1]
    ws.row_dimensions[row].height = 32


def corpo(ws, first, last, ncol, inputs, wrap=False):
    for row in ws.iter_rows(min_row=first, max_row=last, max_col=ncol):
        for cell in row:
            cell.font = Font(name=F, size=9)
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=wrap)
            cell.fill = IN_FILL if cell.column in inputs else LOCK_FILL


def titulo(ws, t, sub):
    ws['A1'] = t
    ws['A1'].font = Font(name=F, size=12, bold=True)
    ws['A2'] = sub
    ws['A2'].font = Font(name=F, size=9, italic=True, color='5B5B63')


def instrucoes(wb, titulo_, linhas):
    ws = wb.active
    ws.title = 'INSTRUÇÕES'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 95
    ws.cell(2, 2, titulo_).font = Font(name=F, size=16, bold=True)
    ws.cell(3, 2, f'Produtos e EANs puxados do cadastro da New Pet na Amazon em {DATA_CADASTRO}.').font = Font(name=F, size=10, color='5B5B63')
    r = 5
    for item in linhas:
        if item[0] == 'h':
            ws.cell(r, 2, item[1]).font = Font(name=F, size=12, bold=True, color=ACC)
        elif item[0] == '':
            pass
        else:
            a = ws.cell(r, 2, item[1]); b = ws.cell(r, 3, item[2])
            a.font = Font(name=F, size=10, bold=True); b.font = Font(name=F, size=10)
            a.alignment = Alignment(vertical='top'); b.alignment = Alignment(wrap_text=True, vertical='top')
            if item[0] == 'in':
                a.fill = IN_FILL
            if item[0] == 'lk':
                a.fill = LOCK_FILL
            ws.row_dimensions[r].height = 30
        r += 1


AVISO_PALITO = ('Palito 8" Rígido 10 un.', 'Cadastrado na Amazon com código de 14 dígitos (17908591301158). Usar esse código no nome dos arquivos, exatamente como está na aba ARQUIVOS.')


# ------------------------------------------------------------------ planilha 1: anúncio
def planilha_anuncio(rows, out):
    wb = Workbook()
    instrucoes(wb, 'PLANILHA DE CONTROLE · IMAGENS DO ANÚNCIO · NEW PET', [
        ('h', 'Como usar'),
        ('k', '1. Pasta no Drive', 'Coloque esta planilha dentro da pasta "NEWPET - IMAGENS ANUNCIO" no Google Drive e abra com Google Planilhas. Na mesma pasta ficam as imagens, sem subpastas.'),
        ('k', '2. Aba ARQUIVOS', 'Lista exata das imagens de cada produto, com o nome pronto. Ao salvar a imagem, copie o nome da coluna "Nome exato do arquivo". Quando ela estiver no Drive, marque o Status como "Na pasta".'),
        ('k', '3. Aba PRODUTOS', 'Mostra o andamento de cada produto. "Completo?" fica SIM quando as 6 imagens obrigatórias estão na pasta. As 3 recomendadas aparecem em coluna própria.'),
        ('k', '4. Aba O QUE MOSTRAR', 'O que mostrar em cada posição, para ração e para petiscos e higiene. Apenas consulta.'),
        ('k', '5. Avise a START', 'Quando os produtos do lote estiverem com "Completo?" = SIM.'),
        ('', ''),
        ('h', 'Cores'),
        ('in', 'Amarelo', 'Vocês preenchem ou escolhem (Status e Observações).'),
        ('lk', 'Cinza', 'Dados do cadastro e fórmulas. Não alterar: são eles que garantem que a imagem vai para o produto certo.'),
        ('', ''),
        ('h', 'Regras dos arquivos'),
        ('k', 'Nome', 'EAN.CÓDIGO.jpg  →  exemplo: 7898968541504.MAIN.jpg. Use exatamente o nome da aba ARQUIVOS: sem sabor, peso, "final", "(1)" ou "cópia".'),
        ('k', 'Imagem', 'JPG, 2000 × 2000 px, cores RGB, até 10 MB.'),
        ('k', 'Textos', 'Não é preciso enviar textos. Título, bullets e descrição são criados pela START. De vocês, só o texto desenhado dentro das imagens.'),
        ('', ''),
        ('h', 'Atenção'),
        ('k', *AVISO_PALITO),
        ('k', 'Produto faltando?', 'Se algum produto à venda não aparece aqui, ou um EAN não corresponde à embalagem, avise a START antes de produzir.'),
    ])

    # O QUE MOSTRAR
    wm = wb.create_sheet('O QUE MOSTRAR')
    titulo(wm, 'O que mostrar em cada imagem do anúncio', 'Consulta. As imagens de exemplo estão no Guia de Imagens do Anúncio.')
    cabecalho(wm, HR, ['Código', 'Onde entra', 'O que mostrar · Ração', 'O que mostrar · Petiscos e higiene', 'Tamanho (px)', 'Obrigatória?'],
              [9, 18, 62, 62, 14, 14])
    for i, g in enumerate(GALERIA):
        for c, v in enumerate([g[0], g[1], g[2], g[3], '2000 × 2000', g[4]], 1):
            wm.cell(HR + 1 + i, c, v)
    ML = HR + len(GALERIA)
    corpo(wm, HR + 1, ML, 6, set(), wrap=True)
    MR = f"'O QUE MOSTRAR'!$A${HR+1}:$A${ML}"

    # PRODUTOS
    wp = wb.create_sheet('PRODUTOS', 1)
    titulo(wp, f'PRODUTOS · New Pet · uma linha por produto à venda (cadastro Amazon em {DATA_CADASTRO})',
           'Cinza = cadastro e fórmulas, não alterar. O andamento é atualizado sozinho a partir do Status da aba ARQUIVOS.')
    cabecalho(wp, HR, ['Nº', 'EAN', 'SKU', 'ASIN', 'Produto (cadastro Amazon)', 'Tipo', 'Peso / tamanho',
                       'Obrigatórias na pasta (de 6)', 'Recomendadas na pasta (de 3)', 'Completo?', 'Observações'],
              [5, 17, 15, 13, 64, 9, 13, 14, 14, 11, 44])
    P0, PL = HR + 1, HR + len(rows)
    for i, r in enumerate(rows):
        rr = P0 + i
        for c, v in enumerate([i + 1, r['ean'], r['sku'], r['asin'], r['nome'], r['tipo'], r['peso']], 1):
            wp.cell(rr, c, v)
        wp.cell(rr, 2).number_format = '@'
        wp.cell(rr, 8, f'=COUNTIFS(ARQUIVOS!$A:$A,$B{rr},ARQUIVOS!$I:$I,"SIM",ARQUIVOS!$J:$J,"Na pasta")')
        wp.cell(rr, 9, f'=COUNTIFS(ARQUIVOS!$A:$A,$B{rr},ARQUIVOS!$I:$I,"Recomendada",ARQUIVOS!$J:$J,"Na pasta")')
        wp.cell(rr, 10, f'=IF(H{rr}>=COUNTIFS(ARQUIVOS!$A:$A,$B{rr},ARQUIVOS!$I:$I,"SIM"),"SIM","NÃO")')
        if r['ean'] == '17908591301158':
            wp.cell(rr, 11, 'Código de 14 dígitos no cadastro: usar exatamente este nos arquivos')
    corpo(wp, P0, PL, 11, {11})
    for rr in range(P0, PL + 1):
        wp.cell(rr, 2).font = Font(name='Courier New', size=9, bold=True)
    wp.conditional_formatting.add(f'J{P0}:J{PL}', FormulaRule(formula=[f'$J{P0}="SIM"'], fill=GREEN, font=Font(name=F, bold=True, color='1F7A3F')))
    wp.conditional_formatting.add(f'J{P0}:J{PL}', FormulaRule(formula=[f'$J{P0}="NÃO"'], fill=RED, font=Font(name=F, color='B3261E')))

    # ARQUIVOS
    wa = wb.create_sheet('ARQUIVOS', 2)
    titulo(wa, 'ARQUIVOS · lista exata de imagens do anúncio por produto',
           'Copie o nome da coluna E ao salvar a imagem e marque o Status "Na pasta" quando ela estiver no Drive.')
    cabecalho(wa, HR, ['EAN', 'Produto', 'Peso / tamanho', 'Código', 'Nome exato do arquivo', 'Onde entra', 'O que mostrar',
                       'Tamanho (px)', 'Obrigatória?', 'Status', 'Observações'],
              [17, 52, 13, 8, 28, 17, 62, 13, 13, 12, 30])
    PR = f'PRODUTOS!$B${P0}:$B${PL}'
    rr = HR + 1
    for r in rows:
        for g in GALERIA:
            m = f'MATCH($A{rr},{PR},0)'
            mm = f'MATCH($D{rr},{MR},0)'
            om = lambda col: f"INDEX('O QUE MOSTRAR'!${col}${HR+1}:${col}${ML},{mm})"
            wa.cell(rr, 1, r['ean']).number_format = '@'
            wa.cell(rr, 2, f'=INDEX(PRODUTOS!$E${P0}:$E${PL},{m})')
            wa.cell(rr, 3, f'=INDEX(PRODUTOS!$G${P0}:$G${PL},{m})')
            wa.cell(rr, 4, g[0])
            wa.cell(rr, 5, f'=$A{rr}&"."&$D{rr}&".jpg"')
            wa.cell(rr, 6, '=' + om('B'))
            wa.cell(rr, 7, f'=IF(INDEX(PRODUTOS!$F${P0}:$F${PL},{m})="Ração",{om("C")},{om("D")})')
            wa.cell(rr, 8, '=' + om('E'))
            wa.cell(rr, 9, '=' + om('F'))
            wa.cell(rr, 10, 'Pendente')
            rr += 1
    AL = rr - 1
    corpo(wa, HR + 1, AL, 11, {10, 11})
    for x in range(HR + 1, AL + 1):
        wa.cell(x, 5).font = Font(name='Courier New', size=9, bold=True)
    dv = DataValidation(type='list', formula1='"Pendente,Na pasta"', allow_blank=False)
    wa.add_data_validation(dv); dv.add(f'J{HR+1}:J{AL}')
    wa.conditional_formatting.add(f'J{HR+1}:J{AL}', FormulaRule(formula=[f'$J{HR+1}="Na pasta"'], fill=GREEN))
    wa.conditional_formatting.add(f'I{HR+1}:I{AL}', FormulaRule(formula=[f'$I{HR+1}="Recomendada"'], font=Font(name=F, color='8A5A00')))

    for ws, last, ncol in ((wp, PL, 11), (wa, AL, 11)):
        ws.auto_filter.ref = f'A{HR}:{chr(64+ncol)}{last}'
        ws.freeze_panes = ws.cell(HR + 1, 3)
    wb.active = 0
    wb.save(out)
    return len(rows), AL - HR


# ------------------------------------------------------------------ planilha 2: A+
def planilha_aplus(rows, out):
    wb = Workbook()
    instrucoes(wb, 'PLANILHA DE CONTROLE · IMAGENS DO A+ · NEW PET', [
        ('h', 'Como usar'),
        ('k', '1. Pasta no Drive', 'Coloque esta planilha dentro da pasta "NEWPET - IMAGENS A+" no Google Drive (pasta separada da do anúncio) e abra com Google Planilhas. Na mesma pasta ficam as imagens, sem subpastas.'),
        ('k', '2. Aba PRODUTOS', 'Escolha o Modelo de A+ (A ou B) e confira a coluna "A+ igual ao EAN". Ela já vem sugerida: os tamanhos de um mesmo produto usam o A+ do maior tamanho. Apague o EAN da coluna para um produto ter A+ próprio.'),
        ('k', '3. Aba ARQUIVOS', 'Lista exata das imagens de cada produto, com o nome pronto. Ela muda sozinha conforme o modelo escolhido e o A+ compartilhado. Linhas com "NÃO" em "Necessário?" não precisam de arquivo.'),
        ('k', '4. Aba MODELOS', 'O que mostrar e o tamanho de cada imagem nos modelos A e B. Apenas consulta.'),
        ('k', '5. Avise a START', 'Quando os produtos do lote estiverem com "Completo?" = SIM.'),
        ('', ''),
        ('h', 'Cores'),
        ('in', 'Amarelo', 'Vocês preenchem ou escolhem (Modelo de A+, A+ igual ao EAN, Status e Observações).'),
        ('lk', 'Cinza', 'Dados do cadastro e fórmulas. Não alterar.'),
        ('', ''),
        ('h', 'Regras dos arquivos'),
        ('k', 'Nome', 'EAN.CÓDIGO.jpg  →  exemplo: 7898968541504.AP01.jpg. Use exatamente o nome da aba ARQUIVOS.'),
        ('k', 'Imagem', 'JPG ou PNG, cores RGB, no tamanho exato indicado (ou maior, na mesma proporção).'),
        ('k', 'Textos', 'Títulos, parágrafos e descrições do A+ são criados pela START. De vocês, só as imagens (com, no máximo, um título curto desenhado).'),
        ('', ''),
        ('h', 'Atenção'),
        ('k', *AVISO_PALITO),
    ])

    wm = wb.create_sheet('MODELOS')
    titulo(wm, 'Modelos de A+ · o que mostrar em cada imagem', 'Consulta. Detalhes e desenhos dos módulos no Guia do Conteúdo A+.')
    cabecalho(wm, HR, ['Código', 'Modelo A · Onde entra', 'Modelo A · O que mostrar (Ração)', 'Modelo A · O que mostrar (Petiscos e higiene)', 'Modelo A · Tamanho (px)',
                       'Modelo B · Onde entra', 'Modelo B · O que mostrar (Ração)', 'Modelo B · O que mostrar (Petiscos e higiene)', 'Modelo B · Tamanho (px)'],
              [9, 26, 44, 44, 13, 30, 40, 40, 13])
    for i, (cod, a, b) in enumerate(APLUS):
        for c, v in enumerate([cod, *a, *b], 1):
            wm.cell(HR + 1 + i, c, v)
    ML = HR + len(APLUS)
    corpo(wm, HR + 1, ML, 9, set(), wrap=True)
    MR = f'MODELOS!$A${HR+1}:$A${ML}'

    wp = wb.create_sheet('PRODUTOS', 1)
    titulo(wp, f'PRODUTOS · New Pet · uma linha por produto à venda (cadastro Amazon em {DATA_CADASTRO})',
           'Amarelo = escolher/preencher · Cinza = cadastro e fórmulas. "A+ igual ao EAN" já vem sugerido; apague para o produto ter A+ próprio.')
    cabecalho(wp, HR, ['Nº', 'EAN', 'SKU', 'ASIN', 'Produto (cadastro Amazon)', 'Tipo', 'Família (mesmo produto, outros tamanhos)', 'Peso / tamanho',
                       'Modelo de A+', 'A+ igual ao EAN', 'Imagens necessárias', 'Imagens na pasta', 'Completo?', 'Observações'],
              [5, 17, 15, 13, 60, 9, 44, 13, 12, 17, 12, 12, 11, 40])
    P0, PL = HR + 1, HR + len(rows)
    for i, r in enumerate(rows):
        rr = P0 + i
        for c, v in enumerate([i + 1, r['ean'], r['sku'], r['asin'], r['nome'], r['tipo'], r['familia'], r['peso'], 'A', r['ref'] or None], 1):
            wp.cell(rr, c, v)
        wp.cell(rr, 2).number_format = '@'
        wp.cell(rr, 10).number_format = '@'
        wp.cell(rr, 11, f'=COUNTIFS(ARQUIVOS!$A:$A,$B{rr},ARQUIVOS!$I:$I,"SIM")')
        wp.cell(rr, 12, f'=COUNTIFS(ARQUIVOS!$A:$A,$B{rr},ARQUIVOS!$I:$I,"SIM",ARQUIVOS!$J:$J,"Na pasta")')
        wp.cell(rr, 13, f'=IF(L{rr}>=K{rr},"SIM","NÃO")')
        if r['ean'] == '17908591301158':
            wp.cell(rr, 14, 'Código de 14 dígitos no cadastro: usar exatamente este nos arquivos')
    corpo(wp, P0, PL, 14, {9, 10, 14})
    for rr in range(P0, PL + 1):
        wp.cell(rr, 2).font = Font(name='Courier New', size=9, bold=True)
    wp.cell(HR, 9).comment = Comment('A = Essencial (7 imagens + COMP). B = Destaques (6 imagens + COMP). Ver Guia do Conteúdo A+.', 'START')
    wp.cell(HR, 10).comment = Comment('EAN do produto cujo A+ este vai repetir. Vazio = A+ próprio (precisa das imagens AP).', 'START')
    dv_mod = DataValidation(type='list', formula1='"A,B"', allow_blank=False, showErrorMessage=True, errorTitle='Modelo de A+', error='Escolha A ou B.')
    dv_ref = DataValidation(type='list', formula1=f'=$B${P0}:$B${PL}', allow_blank=True, showErrorMessage=True,
                            errorTitle='A+ igual ao EAN', error='Use um EAN da coluna B (ou deixe vazio).')
    wp.add_data_validation(dv_mod); wp.add_data_validation(dv_ref)
    dv_mod.add(f'I{P0}:I{PL}'); dv_ref.add(f'J{P0}:J{PL}')
    wp.conditional_formatting.add(f'M{P0}:M{PL}', FormulaRule(formula=[f'$M{P0}="SIM"'], fill=GREEN, font=Font(name=F, bold=True, color='1F7A3F')))
    wp.conditional_formatting.add(f'M{P0}:M{PL}', FormulaRule(formula=[f'$M{P0}="NÃO"'], fill=RED, font=Font(name=F, color='B3261E')))

    wa = wb.create_sheet('ARQUIVOS', 2)
    titulo(wa, 'ARQUIVOS · lista exata de imagens do A+ por produto',
           'Linhas com "NÃO" em Necessário? não precisam de arquivo (o produto usa o A+ de outro EAN, ou o código não existe no modelo escolhido).')
    cabecalho(wa, HR, ['EAN', 'Produto', 'Peso / tamanho', 'Código', 'Nome exato do arquivo', 'Onde entra', 'O que mostrar',
                       'Tamanho (px)', 'Necessário?', 'Status', 'Observações'],
              [17, 50, 13, 8, 28, 30, 50, 13, 30, 12, 26])
    PR = f'PRODUTOS!$B${P0}:$B${PL}'
    rr = HR + 1
    for r in rows:
        for cod, _, _ in APLUS:
            m = f'MATCH($A{rr},{PR},0)'
            mm = f'MATCH($D{rr},{MR},0)'
            modelo = f'INDEX(PRODUTOS!$I${P0}:$I${PL},{m})'
            ref = f'INDEX(PRODUTOS!$J${P0}:$J${PL},{m})'

            def col(a, b):
                return f'IF({modelo}="B",INDEX(MODELOS!${b}${HR+1}:${b}${ML},{mm}),INDEX(MODELOS!${a}${HR+1}:${a}${ML},{mm}))'
            wa.cell(rr, 1, r['ean']).number_format = '@'
            wa.cell(rr, 2, f'=INDEX(PRODUTOS!$E${P0}:$E${PL},{m})')
            wa.cell(rr, 3, f'=INDEX(PRODUTOS!$H${P0}:$H${PL},{m})')
            wa.cell(rr, 4, cod)
            wa.cell(rr, 5, f'=$A{rr}&"."&$D{rr}&".jpg"')
            wa.cell(rr, 6, '=' + col('B', 'F'))
            wa.cell(rr, 7, f'=IF(INDEX(PRODUTOS!$F${P0}:$F${PL},{m})="Ração",{col("C", "G")},{col("D", "H")})')
            wa.cell(rr, 8, '=' + col('E', 'I'))
            # COMP é de cada EAN (cada pacote mostra o próprio peso); AP segue o A+ compartilhado
            wa.cell(rr, 9, f'=IF($D{rr}="COMP","SIM",IF({ref}<>"","NÃO (usa o A+ do EAN "&{ref}&")",IF(AND({modelo}="B",$D{rr}="AP07"),"NÃO (modelo B)","SIM")))')
            wa.cell(rr, 10, 'Pendente')
            rr += 1
    AL = rr - 1
    corpo(wa, HR + 1, AL, 11, {10, 11})
    for x in range(HR + 1, AL + 1):
        wa.cell(x, 5).font = Font(name='Courier New', size=9, bold=True)
    dv = DataValidation(type='list', formula1='"Pendente,Na pasta"', allow_blank=False)
    wa.add_data_validation(dv); dv.add(f'J{HR+1}:J{AL}')
    wa.conditional_formatting.add(f'A{HR+1}:K{AL}', FormulaRule(formula=[f'LEFT($I{HR+1},3)="NÃO"'], font=Font(name=F, color='9A9AA2')))
    wa.conditional_formatting.add(f'J{HR+1}:J{AL}', FormulaRule(formula=[f'$J{HR+1}="Na pasta"'], fill=GREEN))

    for ws, last, ncol in ((wp, PL, 14), (wa, AL, 11)):
        ws.auto_filter.ref = f'A{HR}:{chr(64+ncol)}{last}'
        ws.freeze_panes = ws.cell(HR + 1, 3)
    wb.active = 0
    wb.save(out)
    return len(rows), AL - HR


if __name__ == '__main__':
    rows = preparar(json.load(open(sys.argv[1], encoding='utf-8')))
    print('anúncio', planilha_anuncio([dict(r) for r in rows], sys.argv[2]))
    print('A+', planilha_aplus([dict(r) for r in rows], sys.argv[3]))
