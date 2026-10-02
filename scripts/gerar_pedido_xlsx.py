# ==================================================================
# gerar_pedido_xlsx.py
# Gera um .xlsx por PEDIDO DE COMPRA (PO), no layout de referencia que
# o parceiro usa (logo a esquerda, logo da marca a direita, caixa de
# RESUMO, tabela de itens, rodape) -- so que com o logo da START no
# lugar do logo do parceiro, e o logo da marca/conta no lugar do logo
# do fornecedor. Cores seguem a identidade do logo novo (preto + laranja).
#
# Usado por notificar_pedidos.py (um .xlsx por PO novo, anexado no
# email). Tambem pode ser chamado direto pra gerar um arquivo avulso.
#
# Campos que o modelo de referencia tem e a gente NAO tem (CNPJ, codigo
# de fornecedor, forma/termos de pagamento) ficam de fora -- preferi
# omitir a linha a inventar um valor.
# ==================================================================

import os, json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter

RAIZ = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))
LOGO_START = os.path.join(RAIZ, 'assets', 'start-s-logo.png')

# mesmo mapeamento conta -> logo(s) de assets/logos/ usado no dashboard (app.js, LOGOS).
# Pega sempre o primeiro da lista (logo "principal" da conta); contas sem entrada ficam
# sem logo de marca (so o nome por extenso, em texto).
LOGO_MARCA_CONTA = {
    'alfa_jf':   'treeliss.png',
    'blidshop':  'blidshop.png',
    'conta3':    'orba.png',
    'ozitp':     'kastking.png',
    'jolitex':   'jolitex.png',
    'balboa':    'ligga.png',
    'riomaster': 'riomaster.png',
}

STATUS_PT = {'OPEN': 'Aberto', 'CLOSED': 'Fechado', 'SEM_STATUS': 'Sem status'}

PRETO = '1A1A1A'
LARANJA = 'FD984D'
CINZA_CLARO = 'F7F7F7'
CINZA_BORDA = 'D9D9D9'
FMT_MOEDA = '#,##0.00'


def _fmt_data(iso):
    if not iso:
        return '—'
    return iso[:10][8:10] + '/' + iso[5:7] + '/' + iso[:4]


def _nome_produto(catalogo, asin):
    return ((catalogo or {}).get(asin) or {}).get('nome') or asin


def _add_logo(ws, caminho, celula, altura_px):
    if not os.path.exists(caminho):
        return False
    img = XLImage(caminho)
    img.height, img.width = altura_px, altura_px * (img.width / img.height)
    ws.add_image(img, celula)
    return True


def gerar_xlsx_pedido(conta_chave, conta_nome, po, catalogo, caminho_saida):
    """po = um registro do bloco `pedidos` do dados_vendor.json (um PO, com `itens`)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"PO {po['po']}"[:31]
    ws.sheet_view.showGridLines = False

    bold = Font(bold=True)
    label_cinza = Font(color='595959')
    bold_branco = Font(bold=True, color='FFFFFF')
    titulo_marca = Font(bold=True, size=13, color=PRETO)
    fill_cabecalho_tabela = PatternFill('solid', fgColor=PRETO)
    fill_cinza = PatternFill('solid', fgColor=CINZA_CLARO)
    fill_resumo = PatternFill('solid', fgColor='FBE4D1')  # laranja bem claro
    borda_fina = Border(*[Side(style='thin', color=CINZA_BORDA)] * 4)
    borda_grossa_topo = Border(top=Side(style='medium', color=PRETO))

    # --- linha 1-2: logos (START a esquerda, marca a direita) ---
    ws.row_dimensions[1].height = 46
    _add_logo(ws, LOGO_START, 'A1', 44)
    nome_logo_marca = LOGO_MARCA_CONTA.get(conta_chave)
    caminho_logo_marca = os.path.join(RAIZ, 'assets', 'logos', nome_logo_marca) if nome_logo_marca else None
    if not (caminho_logo_marca and _add_logo(ws, caminho_logo_marca, 'E1', 40)):
        ws['E1'] = conta_nome
        ws['E1'].font = titulo_marca

    # --- cabecalho: dados do pedido (esquerda, colunas A/B) + RESUMO (direita, colunas D/E) ---
    total_itens = len(po.get('itens') or [])
    total_un = (po.get('tot') or {}).get('pedido', 0)
    total_valor = sum((i.get('pedido') or 0) * (i.get('custoUn') or 0) for i in (po.get('itens') or []))
    janela_txt = (f"{_fmt_data(po.get('janelaIni'))} a {_fmt_data(po.get('janelaFim'))}"
                  if po.get('janelaFim') else '—')

    r0 = 3
    ws.cell(row=r0, column=1, value='PEDIDO DE COMPRA').font = Font(bold=True, size=13, color=PRETO)
    ws.cell(row=r0, column=2, value=po['po']).font = Font(bold=True, size=13, color=LARANJA)
    ws.cell(row=r0, column=4, value='RESUMO').font = Font(bold=True, size=11, color=PRETO)
    ws.merge_cells(start_row=r0, start_column=4, end_row=r0, end_column=5)
    ws.cell(row=r0, column=4).fill = fill_resumo
    ws.cell(row=r0, column=5).fill = fill_resumo
    ws.cell(row=r0, column=4).alignment = Alignment(horizontal='center')

    linhas_esquerda = [
        ('Conta', conta_nome),
        ('Centro de distribuição', po.get('destino') or '—'),
        ('Data do pedido', _fmt_data(po.get('data'))),
        ('Status', STATUS_PT.get(po.get('status'), po.get('status') or '—')),
    ]
    linhas_resumo = [
        ('Itens', total_itens, None),
        ('Unidades', total_un, '#,##0'),
        ('Valor solicitado', total_valor, '"R$" ' + FMT_MOEDA),
    ]
    for i, (label, valor) in enumerate(linhas_esquerda):
        rr = r0 + 1 + i
        ws.cell(row=rr, column=1, value=label).font = label_cinza
        ws.cell(row=rr, column=2, value=valor).font = bold
    for i, (label, valor, fmt) in enumerate(linhas_resumo):
        rr = r0 + 1 + i
        ws.cell(row=rr, column=4, value=label).font = label_cinza
        ws.cell(row=rr, column=4).fill = fill_resumo
        c = ws.cell(row=rr, column=5, value=valor)
        c.font = bold
        c.fill = fill_resumo
        c.alignment = Alignment(horizontal='right')
        if fmt:
            c.number_format = fmt

    r = r0 + 1 + max(len(linhas_esquerda), len(linhas_resumo)) + 1

    # Janela de entrega -- destacada (e a data mais acionavel pro cliente: ate quando
    # a Amazon aceita a entrega desse pedido), em faixa laranja de ponta a ponta.
    ws.cell(row=r, column=1, value='JANELA DE ENTREGA').font = Font(bold=True, color='FFFFFF', size=11)
    ws.cell(row=r, column=2, value=janela_txt).font = Font(bold=True, color='FFFFFF', size=12)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=5)
    for col in range(1, 6):
        ws.cell(row=r, column=col).fill = PatternFill('solid', fgColor=LARANJA)
    ws.row_dimensions[r].height = 22
    r += 2

    # --- tabela de itens ---
    linha_cabecalho_tabela = r
    cabecalho = ['ASIN', 'EAN', 'Produto', 'Qtd solicitada', 'Custo unitário (R$)', 'Custo total (R$)']
    for col, titulo in enumerate(cabecalho, start=1):
        c = ws.cell(row=r, column=col, value=titulo)
        c.font = bold_branco
        c.fill = fill_cabecalho_tabela
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 30
    r += 1
    for item in (po.get('itens') or []):
        qtd = item.get('pedido') or 0
        custo_un = item.get('custoUn') or 0
        ws.cell(row=r, column=1, value=item.get('asin'))
        ws.cell(row=r, column=2, value=item.get('ean') or '')
        cel_prod = ws.cell(row=r, column=3, value=_nome_produto(catalogo, item.get('asin')))
        cel_prod.alignment = Alignment(wrap_text=True, vertical='center')
        c_qtd = ws.cell(row=r, column=4, value=qtd)
        c_qtd.alignment = Alignment(horizontal='center')
        c_un = ws.cell(row=r, column=5, value=round(custo_un, 2))
        c_un.number_format = FMT_MOEDA
        c_tot = ws.cell(row=r, column=6, value=round(qtd * custo_un, 2))
        c_tot.number_format = FMT_MOEDA
        if (r - linha_cabecalho_tabela) % 2 == 0:
            for col in range(1, 7):
                ws.cell(row=r, column=col).fill = fill_cinza
        ws.row_dimensions[r].height = 28
        r += 1

    linha_total = r
    c_label = ws.cell(row=r, column=1,
                       value=f'TOTAL — {total_itens} item(ns) · {total_un:,.0f} unidade(s)'.replace(',', '.'))
    c_label.font = bold_branco
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for col in range(1, 6):
        ws.cell(row=r, column=col).fill = fill_cabecalho_tabela
    c_tot_geral = ws.cell(row=r, column=6, value=round(total_valor, 2))
    c_tot_geral.font = bold_branco
    c_tot_geral.fill = fill_cabecalho_tabela
    c_tot_geral.number_format = FMT_MOEDA

    for row in ws.iter_rows(min_row=linha_cabecalho_tabela, max_row=linha_total, min_col=1, max_col=6):
        for cell in row:
            cell.border = borda_fina

    # --- rodape ---
    r += 2
    ws.cell(row=r, column=1, value='Dúvidas ou problemas entre em contato com: comercial@onsetsell.com').font = label_cinza
    r += 1
    ws.cell(row=r, column=1, value='START Inteligência Vendor Central').font = Font(bold=True, color=LARANJA, size=11)
    ws.cell(row=r, column=1).border = borda_grossa_topo

    larguras = {'A': 16, 'B': 18, 'C': 50, 'D': 14, 'E': 18, 'F': 18}
    for col, largura in larguras.items():
        ws.column_dimensions[col].width = largura
    ws.freeze_panes = ws.cell(row=linha_cabecalho_tabela + 1, column=1).coordinate

    os.makedirs(os.path.dirname(caminho_saida), exist_ok=True)
    wb.save(caminho_saida)
    return caminho_saida


if __name__ == '__main__':
    import sys
    if len(sys.argv) != 3:
        print('Uso: python gerar_pedido_xlsx.py <conta> <numero_do_po>')
        sys.exit(1)
    chave, numero_po = sys.argv[1], sys.argv[2]
    contas = json.load(open(os.path.join(RAIZ, 'dados_vendor.json'), encoding='utf-8'))
    conta = contas[chave]
    po = next(p for p in conta['pedidos'] if p['po'] == numero_po)
    saida = os.path.join('/tmp', f'Pedido_{numero_po}_{chave}.xlsx')
    gerar_xlsx_pedido(chave, conta.get('nome') or chave, po, conta.get('catalogo'), saida)
    print(f'Salvo: {saida}')
