# ==================================================================
# classificar_correcoes.py
# Le qualidade_listings.json de todas as contas e separa cada issue em:
#   - AUTO_ATRIBUTO        -- corrigivel enviando um valor de atributo certo
#                             via Listings Items API (patchListingsItem),
#                             desde que a gente saiba o valor correto.
#   - PRECISA_IMAGEM       -- precisa de foto/arquivo de imagem novo ou
#                             hospedagem confiavel; nao da pra corrigir so
#                             com codigo.
#   - PRECISA_DOCUMENTO    -- precisa anexar documento de conformidade
#                             (ficha de seguranca, certificado etc.).
#   - PRECISA_AUTORIZACAO_MARCA -- usa marca/logo protegido sem autorizacao.
#   - PRECISA_SUPORTE_VARIACAO  -- problema de familia de variacao (ASIN
#                             pai/filho); geralmente precisa reestruturar
#                             ou abrir caso com a Amazon.
#   - PRECISA_SUPORTE_AMAZON    -- restricao especifica de caso, so
#                             resolve abrindo caso no Seller/Vendor Central.
#   - OUTRO                -- codigo nao mapeado ainda; revisar manualmente.
#
# NAO ENVIA NADA para a Amazon -- so le o que ja foi capturado por
# capturar_qualidade_listings.py e gera um relatorio pra revisao humana
# (planilha .xlsx + JSON). O envio de correcoes fica pra uma proxima etapa,
# so depois de alguem aprovar o que sera corrigido.
#
# Uso: python scripts/classificar_correcoes.py
# Gera: correcoes_propostas.xlsx e correcoes_propostas.json na raiz do repo.
# ==================================================================

import os, json, glob
import openpyxl
from openpyxl.styles import Font, PatternFill

BASE_DRIVE = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir, 'dados_raw'))
RAIZ = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), os.pardir))

CONTA_NOME = {
    'alfa_jf': 'ALFA JF', 'blidshop': 'Blid Shop', 'petclean': 'Petclean BR',
    'ozitp': 'OZITP', 'jolitex': 'Jolitex', 'balboa': 'Balboa', 'riomaster': 'BR - Rio Master',
}

# codigo da Amazon -> categoria de correcao
CATEGORIAS = {
    # corrigivel enviando atributo certo (precisa do valor correto, mas o mecanismo e simples)
    '18448': 'AUTO_ATRIBUTO',   # atributos principais faltando
    '18002': 'AUTO_ATRIBUTO',   # valor de atributo invalido/incompleto (ex.: preco)
    '99022': 'AUTO_ATRIBUTO',   # value_with_tax sem valores suficientes
    '90244': 'AUTO_ATRIBUTO',   # valor fora da lista aprovada pra esse atributo
    '99016': 'AUTO_ATRIBUTO',   # atributo repetido mais vezes do que o permitido
    '100893': 'AUTO_ATRIBUTO',  # atributo obrigatorio faltando no ASIN filho (ex.: size)
    '18367': 'AUTO_ATRIBUTO',   # tipo de produto foi auto-atualizado (so informativo/confirmar)

    # precisa de imagem/arquivo novo -- nao da pra resolver so com codigo
    '20000': 'PRECISA_IMAGEM',   # timeout baixando a midia (servidor de imagem do cliente)
    '18320': 'PRECISA_IMAGEM',   # imagem principal ausente/incorreta
    '18027': 'PRECISA_IMAGEM',   # imagem com texto/logo/marca-dagua nao permitido
    '100581': 'PRECISA_IMAGEM',  # idem (outra redacao)
    '100588': 'PRECISA_IMAGEM',  # fundo branco irregular
    '100236': 'PRECISA_IMAGEM',  # imagem borrada/pixelada
    '20012': 'PRECISA_IMAGEM',   # imagem excede dimensao maxima
    '100560': 'PRECISA_IMAGEM',  # imagem nao representa o produto
    '100239': 'PRECISA_IMAGEM',  # titulo e imagem nao parecem ser do mesmo produto
    '18254': 'PRECISA_IMAGEM',   # produto muito pequeno no quadro
    '100589': 'PRECISA_IMAGEM',  # idem (outra redacao)
    '300060': 'PRECISA_IMAGEM',  # caminho de arquivo local (file://) enviado por engano, nao uma URL

    # precisa de documento de conformidade
    '18616': 'PRECISA_DOCUMENTO',  # produto quimico -- ficha de dados de seguranca (FISPQ/SDS)
    '100525': 'PRECISA_DOCUMENTO', # documentos de categoria regulada (ex.: consumivel pet)

    # marca/logo protegido sem autorizacao
    '18653': 'PRECISA_AUTORIZACAO_MARCA',

    # familia de variacao (ASIN pai/filho)
    '100898': 'PRECISA_SUPORTE_VARIACAO',  # marca inconsistente entre pai e filho
    '18559': 'PRECISA_SUPORTE_VARIACAO',   # suprimido por problema no ASIN pai

    # caso especifico, so resolve com a Amazon
    '100332': 'PRECISA_SUPORTE_AMAZON',
    '100331': 'PRECISA_SUPORTE_AMAZON',
    '100477': 'PRECISA_SUPORTE_AMAZON',
    '100873': 'PRECISA_SUPORTE_AMAZON',   # codigo de barras ja atribuido a outro ASIN
}

DESCRICAO_CATEGORIA = {
    'AUTO_ATRIBUTO': 'Corrigível via API, com o valor certo do atributo',
    'PRECISA_IMAGEM': 'Precisa de foto/arquivo de imagem novo ou hospedagem confiável',
    'PRECISA_DOCUMENTO': 'Precisa anexar documento de conformidade',
    'PRECISA_AUTORIZACAO_MARCA': 'Uso de marca/logo protegido sem autorização',
    'PRECISA_SUPORTE_VARIACAO': 'Problema de variação (ASIN pai/filho) — investigar/reestruturar',
    'PRECISA_SUPORTE_AMAZON': 'Só resolve abrindo caso no Seller/Vendor Central',
    'OUTRO': 'Código ainda não mapeado — revisar manualmente',
}


def catalogo_da_conta(pasta_raw):
    caminho = os.path.join(pasta_raw, 'catalogo.json')
    if os.path.exists(caminho):
        try:
            return json.load(open(caminho, encoding='utf-8'))
        except Exception:
            return {}
    return {}


def coletar():
    linhas = []
    resumo = {}  # conta -> categoria -> contagem
    for caminho in sorted(glob.glob(os.path.join(BASE_DRIVE, '*', 'raw', 'qualidade_listings.json'))):
        pasta_raw = os.path.dirname(caminho)
        chave_conta = os.path.basename(os.path.dirname(pasta_raw))
        nome_conta = CONTA_NOME.get(chave_conta, chave_conta)
        dados = json.load(open(caminho, encoding='utf-8'))
        catalogo = catalogo_da_conta(pasta_raw)
        resumo.setdefault(nome_conta, {})

        for asin, info in dados.items():
            sku = info.get('sku', '')
            nome_produto = (catalogo.get(asin) or {}).get('nome', '')
            for issue in info.get('issues', []):
                codigo = str(issue.get('code', ''))
                categoria = CATEGORIAS.get(codigo, 'OUTRO')
                resumo[nome_conta][categoria] = resumo[nome_conta].get(categoria, 0) + 1
                linhas.append({
                    'conta': nome_conta,
                    'asin': asin,
                    'sku': sku,
                    'produto': nome_produto,
                    'severidade': issue.get('severity', ''),
                    'codigo': codigo,
                    'categoria': categoria,
                    'descricao_categoria': DESCRICAO_CATEGORIA.get(categoria, ''),
                    'mensagem': issue.get('message', ''),
                })
    return linhas, resumo


def gerar_planilha(linhas, resumo, caminho_saida):
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = 'Resumo'
    ws.append(['Conta', 'Categoria', 'O que significa', 'Quantidade'])
    for cel in ws[1]:
        cel.font = Font(bold=True)
    for conta, categorias in resumo.items():
        for categoria, qtd in sorted(categorias.items(), key=lambda x: -x[1]):
            ws.append([conta, categoria, DESCRICAO_CATEGORIA.get(categoria, ''), qtd])
    for col, largura in zip('ABCD', [18, 28, 55, 12]):
        ws.column_dimensions[col].width = largura

    CORES = {
        'AUTO_ATRIBUTO': 'C6EFCE',
        'PRECISA_IMAGEM': 'FFEB9C',
        'PRECISA_DOCUMENTO': 'FFEB9C',
        'PRECISA_AUTORIZACAO_MARCA': 'FFC7CE',
        'PRECISA_SUPORTE_VARIACAO': 'FFC7CE',
        'PRECISA_SUPORTE_AMAZON': 'FFC7CE',
        'OUTRO': 'D9D9D9',
    }

    ws2 = wb.create_sheet('Todas as correções')
    cabecalho = ['Conta', 'ASIN', 'SKU', 'Produto', 'Severidade', 'Código', 'Categoria', 'O que significa', 'Mensagem da Amazon']
    ws2.append(cabecalho)
    for cel in ws2[1]:
        cel.font = Font(bold=True)
    for l in linhas:
        ws2.append([l['conta'], l['asin'], l['sku'], l['produto'], l['severidade'], l['codigo'],
                    l['categoria'], l['descricao_categoria'], l['mensagem']])
        cor = CORES.get(l['categoria'])
        if cor:
            ws2.cell(row=ws2.max_row, column=7).fill = PatternFill('solid', fgColor=cor)
    for col, largura in zip('ABCDEFGHI', [16, 12, 18, 32, 11, 9, 26, 42, 60]):
        ws2.column_dimensions[col].width = largura
    ws2.freeze_panes = 'A2'
    ws2.auto_filter.ref = ws2.dimensions

    # uma aba só com o que já dá pra corrigir via API (prioridade)
    ws3 = wb.create_sheet('Corrigível via API')
    ws3.append(cabecalho + ['Valor correto (preencher)'])
    for cel in ws3[1]:
        cel.font = Font(bold=True)
    for l in linhas:
        if l['categoria'] == 'AUTO_ATRIBUTO':
            ws3.append([l['conta'], l['asin'], l['sku'], l['produto'], l['severidade'], l['codigo'],
                        l['categoria'], l['descricao_categoria'], l['mensagem'], ''])
    for col, largura in zip('ABCDEFGHIJ', [16, 12, 18, 32, 11, 9, 15, 42, 60, 30]):
        ws3.column_dimensions[col].width = largura
    ws3.freeze_panes = 'A2'
    ws3.auto_filter.ref = ws3.dimensions

    wb.save(caminho_saida)


def main():
    linhas, resumo = coletar()
    print(f'Total de issues classificadas: {len(linhas)}\n')
    for conta, categorias in resumo.items():
        print(f'=== {conta} ===')
        for categoria, qtd in sorted(categorias.items(), key=lambda x: -x[1]):
            print(f'  {qtd:4d}  {categoria}')
        print()

    caminho_json = os.path.join(RAIZ, 'correcoes_propostas.json')
    with open(caminho_json, 'w', encoding='utf-8') as f:
        json.dump({'linhas': linhas, 'resumo': resumo}, f, ensure_ascii=False, indent=2)
    print(f'Salvo: {caminho_json}')

    caminho_xlsx = os.path.join(RAIZ, 'correcoes_propostas.xlsx')
    gerar_planilha(linhas, resumo, caminho_xlsx)
    print(f'Salvo: {caminho_xlsx}')


if __name__ == '__main__':
    main()
